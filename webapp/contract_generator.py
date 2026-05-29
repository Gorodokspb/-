"""Contract DOCX generation service.

Generates a contract DOCX from the template by replacing [[PLACEHOLDER]]
markers with project + counterparty + contract settings data.
Uses python-docx for safe template processing on Linux (no Word COM).
"""

from __future__ import annotations

import datetime
import re
import shutil
import subprocess
import tempfile
import uuid
from copy import deepcopy
from pathlib import Path

from docx import Document as DocxDocument

from webapp.config import get_settings
from webapp.storage import sanitize_filename, storage_relative_path, resolve_storage_path


TEMPLATE_FILENAME = "contract_template_physical.docx"


def _get_template_path() -> Path:
    settings = get_settings()
    template = settings.storage_root / TEMPLATE_FILENAME
    if template.exists():
        return template
    base_dir = Path(__file__).resolve().parent.parent
    fallback = base_dir / TEMPLATE_FILENAME
    if fallback.exists():
        return fallback
    raise FileNotFoundError(f"Contract template not found: {template} or {fallback}")


def _get_contracts_dir() -> Path:
    settings = get_settings()
    contracts_dir = settings.contracts_dir
    contracts_dir.mkdir(parents=True, exist_ok=True)
    return contracts_dir


def _normalize_counterparty_type(raw_type: str) -> str:
    t = (raw_type or "").strip().lower()
    if t in ("физлицо", "физическое лицо", "physical", "person", "физ. лицо"):
        return "Физическое лицо"
    if t in ("ип", "индивидуальный предприниматель"):
        return "ИП"
    if t in ("юридическое лицо ооо", "юридическое лицо", "ооо", "юрлицо", "company"):
        return "Юридическое лицо ООО"
    return raw_type or ""


def _counterparty_display_name(row: dict | None) -> str:
    if not row:
        return ""
    full = str(row.get("full_name") or "").strip()
    if full:
        return full
    name = str(row.get("name") or "").strip()
    if name:
        return name
    company = str(row.get("company_name") or "").strip()
    if company:
        return company
    return ""


def _detect_customer_gender(full_name: str) -> str:
    parts = [p for p in str(full_name or "").strip().split() if p]
    if len(parts) >= 3:
        patronymic = parts[2].lower()
        if patronymic.endswith(("вна", "ична", "кызы")):
            return "female"
        if patronymic.endswith(("вич", "оглы")):
            return "male"
    if len(parts) >= 2:
        first_name = parts[1].lower()
        if first_name.endswith(("а", "я")) and not first_name.endswith(("илья", "никита", "кузьма", "фома")):
            return "female"
    return "male" if parts else "unknown"


def pluralize(value: int | float, forms: tuple[str, str, str]) -> str:
    number = abs(int(value)) % 100
    if 11 <= number <= 19:
        return forms[2]
    remainder = number % 10
    if remainder == 1:
        return forms[0]
    if 2 <= remainder <= 4:
        return forms[1]
    return forms[2]


def number_to_words_ru(value: int) -> str:
    units = {
        0: "ноль", 1: "один", 2: "два", 3: "три", 4: "четыре",
        5: "пять", 6: "шесть", 7: "семь", 8: "восемь", 9: "девять",
        10: "десять", 11: "одиннадцать", 12: "двенадцать", 13: "тринадцать",
        14: "четырнадцать", 15: "пятнадцать", 16: "шестнадцать",
        17: "семнадцать", 18: "восемнадцать", 19: "девятнадцать",
    }
    tens = {
        20: "двадцать", 30: "тридцать", 40: "сорок", 50: "пятьдесят",
        60: "шестьдесят", 70: "семьдесят", 80: "восемьдесят", 90: "девяносто",
    }
    hundreds = {
        100: "сто", 200: "двести", 300: "триста", 400: "четыреста",
        500: "пятьсот", 600: "шестьсот", 700: "семьсот", 800: "восемьсот",
        900: "девятьсот",
    }
    groups = [
        ("", "", "", False),
        ("тысяча", "тысячи", "тысяч", True),
        ("миллион", "миллиона", "миллионов", False),
        ("миллиард", "миллиарда", "миллиардов", False),
    ]

    number = int(value)
    if number == 0:
        return units[0]

    parts = []
    group_index = 0
    while number > 0 and group_index < len(groups):
        chunk = number % 1000
        number //= 1000
        if chunk:
            words = []
            h = chunk // 100 * 100
            if h:
                words.append(hundreds[h])
            last_two = chunk % 100
            if last_two < 20:
                if last_two:
                    if last_two == 1:
                        words.append("одна" if groups[group_index][3] else "один")
                    elif last_two == 2:
                        words.append("две" if groups[group_index][3] else "два")
                    else:
                        words.append(units[last_two])
            else:
                t = last_two // 10 * 10
                u = last_two % 10
                words.append(tens[t])
                if u:
                    if u == 1:
                        words.append("одна" if groups[group_index][3] else "один")
                    elif u == 2:
                        words.append("две" if groups[group_index][3] else "два")
                    else:
                        words.append(units[u])
            if group_index > 0:
                words.append(pluralize(chunk, groups[group_index][:3]))
            parts.insert(0, " ".join(words))
        group_index += 1
    return " ".join(parts)


def parse_money_value(raw_value) -> float | None:
    text = str(raw_value or "").strip()
    if not text:
        return None
    cleaned = re.sub(r"[^0-9,.\-]", "", text).replace(",", ".")
    if cleaned.count(".") > 1:
        head, tail = cleaned.rsplit(".", 1)
        cleaned = head.replace(".", "") + "." + tail
    try:
        return round(float(cleaned), 2)
    except ValueError:
        return None


def format_money_with_words(raw_value) -> str:
    amount = parse_money_value(raw_value)
    if amount is None:
        return str(raw_value or "").strip()
    rubles = int(amount)
    kopecks = int(round((amount - rubles) * 100))
    amount_text = f"{rubles:,}".replace(",", " ")
    rubles_words = number_to_words_ru(rubles)
    return (
        f"{amount_text} ({rubles_words}) "
        f"{pluralize(rubles, ('рубль', 'рубля', 'рублей'))} "
        f"{kopecks:02d} {pluralize(kopecks, ('копейка', 'копейки', 'копеек'))}"
    )


def parse_date_value(raw_value: str):
    text = str(raw_value or "").strip()
    if not text:
        return None
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def format_contract_date(raw_value: str) -> str:
    months = {
        1: "января", 2: "февраля", 3: "марта", 4: "апреля",
        5: "мая", 6: "июня", 7: "июля", 8: "августа",
        9: "сентября", 10: "октября", 11: "ноября", 12: "декабря",
    }
    text = (raw_value or "").strip()
    if not text:
        today = datetime.date.today()
        return f'" {today.day:02d} " {months[today.month]} {today.year} г.'
    try:
        dt = datetime.datetime.strptime(text, "%d.%m.%Y").date()
        return f'" {dt.day:02d} " {months[dt.month]} {dt.year} г.'
    except ValueError:
        return text


def format_long_date(raw_value: str, fallback: str = "") -> str:
    months = {
        1: "января", 2: "февраля", 3: "марта", 4: "апреля",
        5: "мая", 6: "июня", 7: "июля", 8: "августа",
        9: "сентября", 10: "октября", 11: "ноября", 12: "декабря",
    }
    dt = parse_date_value(raw_value)
    if dt:
        return f"{dt.day} {months[dt.month]} {dt.year} года"
    return str(raw_value or fallback).strip()


def format_deadline_text(raw_value: str) -> str:
    months = {
        1: "января", 2: "февраля", 3: "марта", 4: "апреля",
        5: "мая", 6: "июня", 7: "июля", 8: "августа",
        9: "сентября", 10: "октября", 11: "ноября", 12: "декабря",
    }
    dt = parse_date_value(raw_value)
    if dt:
        return f"не позднее «{dt.day}» {months[dt.month]} {dt.year} г.;"
    text = str(raw_value or "").strip()
    return text or "не указано;"


def _get_executor_profile(contractor_mode: str) -> dict:
    if contractor_mode == "ip":
        return {"email": "gorodok198@yandex.ru", "label": "ИП"}
    return {"email": "info@dekorartstroy.ru", "label": "ООО"}


def _build_customer_intro(counterparty_row: dict | None, settings: dict) -> str:
    intro_override = str(settings.get("intro_override", "")).strip()
    if intro_override:
        return intro_override
    counterparty_type = _normalize_counterparty_type(counterparty_row.get("type", "")) if counterparty_row else ""
    display_name = str(
        settings.get("customer_name")
        or (_counterparty_display_name(counterparty_row) if counterparty_row else "")
        or "Заказчик"
    ).strip()
    contractor_mode = settings.get("contractor_mode", "ooo")

    if contractor_mode == "ip":
        contractor_intro = (
            "Индивидуальный предприниматель Гордеев Алексей Николаевич, "
            "именуемый в дальнейшем «Подрядчик», с одной стороны, "
        )
    else:
        contractor_intro = (
            "Общество с ограниченной ответственностью «ДекорАртСтрой» именуемое в дальнейшем "
            "«Подрядчик», в лице Генерального директора Шарипова Шехрозжона Шавкатовича, "
            "действующего на основании Устава с одной стороны, "
        )

    if counterparty_type == "Физическое лицо":
        gender = settings.get("customer_gender", "auto")
        if gender == "auto":
            gender = _detect_customer_gender(display_name)
        if gender == "female":
            customer_intro = f"и гражданка {display_name}, именуемая в дальнейшем «Заказчик», "
        elif gender == "male":
            customer_intro = f"и гражданин {display_name}, именуемый в дальнейшем «Заказчик», "
        else:
            customer_intro = f"и гражданин(ка) {display_name}, именуемый(ая) в дальнейшем «Заказчик», "
        return contractor_intro + customer_intro + "вместе именуемые «Стороны», заключили настоящий договор (далее – «Договор») о нижеследующем:"

    if counterparty_type == "Юридическое лицо ООО":
        company_name = (counterparty_row or {}).get("company_name") or display_name
        director_name = (counterparty_row or {}).get("director_name") or "уполномоченного представителя"
        director_basis = (counterparty_row or {}).get("director_basis") or "Устава"
        return (
            contractor_intro
            + f"и {company_name}, именуемое в дальнейшем «Заказчик», в лице {director_name}, "
            f"действующего на основании {director_basis}, "
            "вместе именуемые «Стороны», заключили настоящий договор (далее – «Договор») о нижеследующем:"
        )

    company_name = (counterparty_row or {}).get("company_name") or display_name
    return (
        contractor_intro
        + f"и Индивидуальный предприниматель {company_name}, именуемый в дальнейшем «Заказчик», "
        "вместе именуемые «Стороны», заключили настоящий договор (далее – «Договор») о нижеследующем:"
    )


def _build_customer_contract_clause(counterparty_row: dict | None, settings: dict) -> str:
    intro_override = str(settings.get("intro_override", "")).strip()
    if intro_override:
        return intro_override
    display_name = str(
        settings.get("customer_name")
        or (_counterparty_display_name(counterparty_row) if counterparty_row else "")
        or "Заказчик"
    ).strip()
    gender = settings.get("customer_gender", "auto")
    if gender == "auto":
        gender = _detect_customer_gender(display_name)
    if gender == "female":
        return f"и гражданка {display_name}, именуемая в дальнейшем «Заказчик»,"
    if gender == "male":
        return f"и гражданин {display_name}, именуемый в дальнейшем «Заказчик»,"
    return f"и гражданин(ка) {display_name}, именуемый(ая) в дальнейшем «Заказчик»,"


def _build_dynamic_payment_block(settings: dict) -> str:
    override = str(settings.get("payments_override", "")).strip()
    if override:
        return override.replace("\n", "\r")
    payments = settings.get("payments") or []
    lines = []
    for index, payment in enumerate(payments, start=1):
        payment_date = format_long_date(payment.get("date", ""), fallback=f"дата платежа {index}")
        payment_amount = format_money_with_words(payment.get("amount", "")) or "0 (ноль) рублей 00 копеек"
        lines.append(
            f"4.4.{index}. В срок не позднее «{payment_date}» «Заказчик» выплачивает «Подрядчику» денежные средства "
            f"в размере {payment_amount}, НДС не облагается."
        )
    return "\r".join(lines)


def _build_payment_lines_for_template(settings: dict) -> tuple[str, str, str]:
    dynamic_block = _build_dynamic_payment_block(settings)
    if dynamic_block:
        lines = [line.strip() for line in dynamic_block.split("\r") if line.strip()]
        line_1 = lines[0] if len(lines) > 0 else ""
        line_2 = lines[1] if len(lines) > 1 else ""
        line_3_plus = "\r".join(lines[2:]) if len(lines) > 2 else ""
        return line_1, line_2, line_3_plus
    payments = settings.get("payments") or []
    if not payments:
        return "", "", ""
    line_1 = (
        f"4.4.1. В срок не позднее «{format_long_date(payments[0].get('date', ''))}» "
        f"«Заказчик» выплачивает «Подрядчику» денежные средства "
        f"в размере {format_money_with_words(payments[0].get('amount', '')) or '0 (ноль) рублей 00 копеек'}, НДС не облагается."
    ) if len(payments) > 0 else ""
    line_2 = (
        f"4.4.2. В срок не позднее «{format_long_date(payments[1].get('date', ''))}» "
        f"«Заказчик» выплачивает «Подрядчику» денежные средства "
        f"в размере {format_money_with_words(payments[1].get('amount', '')) or '0 (ноль) рублей 00 копеек'}, НДС не облагается."
    ) if len(payments) > 1 else ""
    tail_lines = []
    for index, payment in enumerate(payments[2:], start=3):
        payment_date = format_long_date(payment.get("date", ""), fallback=f"дата платежа {index}")
        payment_amount = format_money_with_words(payment.get("amount", "")) or "0 (ноль) рублей 00 копеек"
        tail_lines.append(
            f"4.4.{index}. В срок не позднее «{payment_date}» «Заказчик» выплачивает «Подрядчику» денежные средства "
            f"в размере {payment_amount}, НДС не облагается."
        )
    return line_1, line_2, "\r".join(tail_lines)


def _build_communications_block(settings: dict) -> str:
    override = str(settings.get("communications_override", "")).strip()
    if override:
        return override.replace("\n", "\r")
    customer_email = settings.get("customer_email", "").strip() or "не указан"
    contractor_email = settings.get("contractor_email", "").strip() or _get_executor_profile(settings.get("contractor_mode", "ooo"))["email"]
    working_group_text = settings.get("working_group_text", "").strip()
    lines = [
        "8.1. Стороны пришли к соглашению об использовании в рамках настоящего Договора следующих адресов электронной почты:",
        f"Заказчик: {customer_email},",
        f"Подрядчик: {contractor_email} ,",
    ]
    if working_group_text:
        lines.append(working_group_text)
    lines.append(
        "С подписанием настоящего Договора, Стороны признают, что направление уведомления и (или) сообщения, "
        "связанных с исполнением настоящего Договора, по адресам электронной почты указанных в настоящем пункте, "
        "считается надлежащим уведомлением сторон. Датой получения уведомления и (или) сообщения Стороной "
        "настоящего Договора, будет считаться день, следующий за днем направления исходящего письма другой "
        "Стороны с адреса электронной почты указанного в настоящем пункте Договора. При исполнении Договора "
        "Стороны обязуются надлежащим образом обеспечивать режим доступа к указанным адресам электронной почты. "
        "Сторона, не обеспечившая режим доступа к адресу электронной почты, указанной в настоящем пункте Договора, "
        "самостоятельно несет риск наступления неблагоприятных последствий, связанных с исполнением указанной обязанности."
    )
    return "\r".join(lines)


def build_contract_replacements(
    project: dict,
    counterparty: dict | None,
    settings: dict,
    estimate_total: str = "0",
) -> dict[str, str]:
    contract_number = (settings.get("contract_number") or project.get("contract") or "б/н").strip()
    contract_date = format_contract_date(settings.get("contract_date") or project.get("date") or "")
    counterparty_type = _normalize_counterparty_type((counterparty or {}).get("type", "")) if counterparty else ""
    display_name = str(
        settings.get("customer_name")
        or (_counterparty_display_name(counterparty) if counterparty else "")
        or project.get("customer") or "Заказчик"
    ).strip()

    if counterparty and counterparty_type == "Физическое лицо":
        passport_main = settings.get("passport_series_number") or counterparty.get("passport_series_number") or "не указано"
        issued_by = settings.get("passport_issued_by") or counterparty.get("passport_issued_by") or "не указано"
        department_code = settings.get("passport_department_code") or counterparty.get("passport_department_code") or "не указано"
        registration_address = settings.get("registration_address") or counterparty.get("registration_address") or "не указано"
        work_address = settings.get("object_address") or counterparty.get("work_address") or project.get("address") or project.get("project_name") or registration_address
        phone = settings.get("customer_phone") or counterparty.get("phone") or "не указан"
        email = settings.get("customer_email") or counterparty.get("email") or "не указан"
    elif counterparty and counterparty_type == "Юридическое лицо ООО":
        passport_main = counterparty.get("inn") or "не указано"
        issued_by = " / ".join(part for part in [counterparty.get("kpp", ""), counterparty.get("ogrn", "")] if part) or "не указано"
        department_code = " / ".join(part for part in [counterparty.get("bank_name", ""), counterparty.get("bank_bik", "")] if part) or "не указано"
        registration_address = counterparty.get("legal_address") or "не указано"
        work_address = settings.get("object_address") or counterparty.get("work_address") or project.get("address") or project.get("project_name") or registration_address
        phone = settings.get("customer_phone") or counterparty.get("phone") or "не указан"
        email = settings.get("customer_email") or counterparty.get("email") or "не указан"
    elif counterparty:
        passport_main = counterparty.get("inn") or "не указано"
        issued_by = counterparty.get("ogrnip") or "не указано"
        department_code = " / ".join(
            part for part in [
                counterparty.get("bank_name", ""),
                counterparty.get("bank_bik", ""),
            ] if part
        ) or "не указано"
        registration_address = settings.get("registration_address") or counterparty.get("legal_address") or "не указано"
        work_address = settings.get("object_address") or counterparty.get("work_address") or project.get("address") or project.get("project_name") or registration_address
        phone = settings.get("customer_phone") or counterparty.get("phone") or "не указан"
        email = settings.get("customer_email") or counterparty.get("email") or "не указан"
    else:
        passport_main = "не указано"
        issued_by = "не указано"
        department_code = "не указано"
        registration_address = "не указано"
        work_address = settings.get("object_address") or project.get("address") or project.get("project_name") or "не указано"
        phone = "не указан"
        email = "не указан"

    object_address = settings.get("object_address") or (counterparty.get("work_address") if counterparty else "") or project.get("address") or project.get("project_name") or "не указано"
    work_end_date = format_deadline_text(settings.get("work_end_date"))

    price_total_value = parse_money_value(settings.get("price_total"))
    if price_total_value is None:
        try:
            et = float(estimate_total.replace(" ", "").replace(",", "."))
            if et > 0:
                price_total_value = et
        except (ValueError, TypeError, AttributeError):
            pass

    advance_value = parse_money_value(settings.get("advance_amount")) or 0.0
    payment_values = []
    for payment in settings.get("payments") or []:
        payment_amount = parse_money_value(payment.get("amount"))
        if payment_amount is not None:
            payment_values.append(payment_amount)

    computed_final = None
    if price_total_value is not None:
        computed_final = round(price_total_value - advance_value - sum(payment_values), 2)
    final_payment_value = parse_money_value(settings.get("final_payment_amount"))
    if final_payment_value is None:
        final_payment_value = computed_final

    price_total = format_money_with_words(price_total_value) if price_total_value is not None else ""
    advance_amount = format_money_with_words(advance_value)
    final_payment_amount = format_money_with_words(final_payment_value) if final_payment_value is not None else ""

    materials_mode = settings.get("materials_mode", "customer")
    materials_phrase = (
        "выполняются из материалов, предоставляемых «Заказчиком»"
        if materials_mode == "customer"
        else "выполняются с возможностью закупки материалов «Подрядчиком» по согласованию с «Заказчиком»"
    )

    payment_1_line, payment_2_line, payment_3_line = _build_payment_lines_for_template(settings)

    replacements = {
        "[[CONTRACT_NUMBER]]": f"ДОГОВОР № {contract_number}",
        "[[CONTRACT_DATE]]": contract_date,
        "[[CUSTOMER_CLAUSE]]": _build_customer_contract_clause(counterparty, settings),
        "[[CUSTOMER_INTRO]]": _build_customer_intro(counterparty, settings),
        "[[CUSTOMER_NAME]]": display_name,
        "[[OBJECT_ADDRESS]]": object_address,
        "[[PASSPORT]]": passport_main,
        "[[PASSPORT_ISSUED_BY]]": issued_by,
        "[[PASSPORT_CODE]]": department_code,
        "[[REGISTRATION_ADDRESS]]": registration_address,
        "[[WORK_ADDRESS]]": work_address,
        "[[CUSTOMER_PHONE]]": phone,
        "[[CUSTOMER_EMAIL]]": email,
        "[[WORK_END_DATE]]": work_end_date,
        "[[PRICE_TOTAL]]": price_total,
        "[[FINAL_PAYMENT]]": final_payment_amount,
        "[[ADVANCE_PAYMENT_LINE]]": f"В день подписания настоящего Договора «Заказчик» выплачивает «Подрядчику» авансовый платеж в размере {advance_amount}, НДС не облагается.",
        "[[PAYMENT_LINE_1]]": payment_1_line,
        "[[PAYMENT_LINE_2]]": payment_2_line,
        "[[PAYMENT_LINE_3_PLUS]]": payment_3_line,
        "[[PAYMENTS_BLOCK]]": _build_dynamic_payment_block(settings),
        "[[CONTRACTOR_EMAIL]]": settings.get("contractor_email") or _get_executor_profile(settings.get("contractor_mode", "ooo"))["email"],
        "[[WORKING_GROUP_TEXT]]": (settings.get("working_group_text") or "").strip(),
        "[[COMMUNICATIONS_BLOCK]]": _build_communications_block(settings),
    }

    materials_replacement = (
        f"Работы, предусмотренные настоящим Договором, {materials_phrase} (давальческий материал) до начала выполнения Работ."
        if materials_mode == "customer"
        else f"Работы, предусмотренные настоящим Договором, {materials_phrase}."
    )
    replacements["Работы, предусмотренные настоящим Договором, выполняются из материалов, предоставляемых «Заказчиком» (давальческий материал) до начала выполнения Работ."] = materials_replacement

    return {k: v for k, v in replacements.items() if v}


def replace_placeholders_in_docx(template_path: Path, replacements: dict[str, str]) -> DocxDocument:
    doc = DocxDocument(str(template_path))

    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, replacements)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, replacements)

    for section in doc.sections:
        for header in [section.header, section.first_page_header]:
            if header and header.paragraphs:
                for paragraph in header.paragraphs:
                    _replace_in_paragraph(paragraph, replacements)
        for footer in [section.footer, section.first_page_footer]:
            if footer and footer.paragraphs:
                for paragraph in footer.paragraphs:
                    _replace_in_paragraph(paragraph, replacements)

    _replace_in_xml_text_nodes(doc, replacements, _W_NS)

    working_group_text = replacements.get("[[WORKING_GROUP_TEXT]]", "").strip()
    if working_group_text:
        prefixed_text = working_group_text
        if not prefixed_text.lower().startswith("рабочая группа"):
            prefixed_text = f"Рабочая группа: {working_group_text}"
        _replace_working_group_paragraph(doc, prefixed_text)

    replacement_values = [v for v in replacements.values() if v]
    _normalize_replacement_colors(doc, replacement_values)

    _remove_red_colors(doc)

    return doc


_WORKING_GROUP_MARKER = "Рабочая группа WhatsApp"
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _remove_red_colors(doc: DocxDocument):
    from lxml import etree
    body = doc.element.body
    for color_elem in body.iter(f"{{{_W_NS}}}color"):
        val = color_elem.get(f"{{{_W_NS}}}val")
        if val and val.upper() == "FF0000":
            color_elem.set(f"{{{_W_NS}}}val", "000000")
            for attr_name in [f"{{{_W_NS}}}theme", f"{{{_W_NS}}}tint", f"{{{_W_NS}}}shade"]:
                for attr_key in list(color_elem.attrib.keys()):
                    if attr_key == attr_name:
                        del color_elem.attrib[attr_key]
            rPr = color_elem.getparent()
            if rPr is not None:
                for tag in [f"{{{_W_NS}}}themeColor", f"{{{_W_NS}}}themeTint", f"{{{_W_NS}}}themeShade"]:
                    child = rPr.find(tag)
                    if child is not None:
                        rPr.remove(child)


def _normalize_run_color_to_black(run):
    from lxml import etree
    rPr = run._element.find(f"{{{_W_NS}}}rPr")
    if rPr is None:
        rPr = etree.SubElement(run._element, f"{{{_W_NS}}}rPr")
        run._element.insert(0, rPr)
    color_elem = rPr.find(f"{{{_W_NS}}}color")
    if color_elem is None:
        color_elem = etree.SubElement(rPr, f"{{{_W_NS}}}color")
    color_elem.set(f"{{{_W_NS}}}val", "000000")
    for attr_name in [f"{{{_W_NS}}}theme", f"{{{_W_NS}}}tint", f"{{{_W_NS}}}shade"]:
        for attr_key in list(color_elem.attrib.keys()):
            if attr_key == attr_name:
                del color_elem.attrib[attr_key]
    for tag in [f"{{{_W_NS}}}themeColor", f"{{{_W_NS}}}themeTint", f"{{{_W_NS}}}themeShade"]:
        child = rPr.find(tag)
        if child is not None:
            rPr.remove(child)


def _normalize_replacement_colors(doc: DocxDocument, replacement_values: list[str]):
    from lxml import etree
    significant_values = [v for v in replacement_values if v and len(v) >= 3]
    if not significant_values:
        return

    def _text_contains_value(text):
        for val in significant_values:
            if val in text:
                return True
        return False

    def _normalize_run_color_in_element(r_elem):
        rPr = r_elem.find(f"{{{_W_NS}}}rPr")
        if rPr is None:
            rPr = etree.SubElement(r_elem, f"{{{_W_NS}}}rPr")
            r_elem.insert(0, rPr)
        color_elem = rPr.find(f"{{{_W_NS}}}color")
        if color_elem is None:
            color_elem = etree.SubElement(rPr, f"{{{_W_NS}}}color")
        color_elem.set(f"{{{_W_NS}}}val", "000000")
        for attr_name in [f"{{{_W_NS}}}theme", f"{{{_W_NS}}}tint", f"{{{_W_NS}}}shade"]:
            for attr_key in list(color_elem.attrib.keys()):
                if attr_key == attr_name:
                    del color_elem.attrib[attr_key]
        for tag in [f"{{{_W_NS}}}themeColor", f"{{{_W_NS}}}themeTint", f"{{{_W_NS}}}themeShade"]:
            child = rPr.find(tag)
            if child is not None:
                rPr.remove(child)

    def _normalize_paragraph(paragraph):
        for run in paragraph.runs:
            run_text = run.text or ""
            if run_text.strip():
                _normalize_run_color_to_black(run)

    def _normalize_xml_paragraph(p_elem):
        if not _text_contains_value("".join(t.text or "" for t in p_elem.iter(f"{{{_W_NS}}}t"))):
            return
        for r_elem in p_elem.iter(f"{{{_W_NS}}}r"):
            t_elem = r_elem.find(f"{{{_W_NS}}}t")
            if t_elem is not None and (t_elem.text or "").strip():
                _normalize_run_color_in_element(r_elem)

    for paragraph in doc.paragraphs:
        if _text_contains_value(paragraph.text or ""):
            _normalize_paragraph(paragraph)
            _normalize_xml_paragraph(paragraph._element)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if _text_contains_value(paragraph.text or ""):
                        _normalize_paragraph(paragraph)
                        _normalize_xml_paragraph(paragraph._element)


def _replace_working_group_paragraph(doc: DocxDocument, working_group_text: str):
    from lxml import etree

    target_paragraph = None
    for paragraph in doc.paragraphs:
        if _WORKING_GROUP_MARKER in paragraph.text:
            target_paragraph = paragraph
            break

    if target_paragraph is None:
        return

    p_elem = target_paragraph._element

    child_tags_to_remove = [f"{{{_W_NS}}}r", f"{{{_W_NS}}}hyperlink"]
    children_to_remove = []
    for child in p_elem:
        if child.tag in child_tags_to_remove:
            children_to_remove.append(child)

    for child in children_to_remove:
        p_elem.remove(child)

    new_run = etree.SubElement(p_elem, f"{{{_W_NS}}}r")
    new_rPr = etree.SubElement(new_run, f"{{{_W_NS}}}rPr")
    new_rFonts = etree.SubElement(new_rPr, f"{{{_W_NS}}}rFonts")
    new_rFonts.set(f"{{{_W_NS}}}ascii", "Times New Roman")
    new_rFonts.set(f"{{{_W_NS}}}hAnsi", "Times New Roman")
    new_sz = etree.SubElement(new_rPr, f"{{{_W_NS}}}sz")
    new_sz.set(f"{{{_W_NS}}}val", "24")
    new_szCs = etree.SubElement(new_rPr, f"{{{_W_NS}}}szCs")
    new_szCs.set(f"{{{_W_NS}}}val", "24")
    new_color = etree.SubElement(new_rPr, f"{{{_W_NS}}}color")
    new_color.set(f"{{{_W_NS}}}val", "000000")
    new_u = etree.SubElement(new_rPr, f"{{{_W_NS}}}u")
    new_u.set(f"{{{_W_NS}}}val", "none")
    etree.SubElement(new_rPr, f"{{{_W_NS}}}noProof")
    new_t = etree.SubElement(new_run, f"{{{_W_NS}}}t")
    new_t.set(f"{{http://www.w3.org/XML/1998/namespace}}space", "preserve")
    new_t.text = working_group_text


def _replace_in_xml_text_nodes(doc: DocxDocument, replacements: dict[str, str], w_ns: str):
    body = doc.element.body
    for t_elem in body.iter("{%s}t" % w_ns):
        if t_elem.text:
            for key, value in replacements.items():
                if key in t_elem.text:
                    t_elem.text = t_elem.text.replace(key, value)


def _replace_in_paragraph(paragraph, replacements: dict[str, str]):
    full_text = paragraph.text
    has_placeholder = any(key in full_text for key in replacements)
    if not has_placeholder:
        return

    for key, value in replacements.items():
        for run in paragraph.runs:
            if key in run.text:
                run.text = run.text.replace(key, value)


def generate_contract_docx(project_id: int) -> dict:
    from webapp.db import (
        fetch_project,
        fetch_counterparty,
        fetch_contract_settings,
        get_project_estimate_total,
        create_or_update_contract_document,
        update_document_file_path,
    )

    project = fetch_project(project_id)
    if not project:
        raise ValueError(f"Project {project_id} not found")

    counterparty = None
    counterparty_id = project.get("counterparty_id")
    if counterparty_id:
        counterparty = fetch_counterparty(int(counterparty_id))

    settings = fetch_contract_settings(project_id)
    estimate_total = get_project_estimate_total(project_id)

    settings.setdefault("contract_number", project.get("contract") or "")
    settings.setdefault("contract_date", project.get("date") or "")
    settings.setdefault("customer_gender", "auto")
    settings.setdefault("contractor_mode", "ooo")
    settings.setdefault("materials_mode", "customer")
    settings.setdefault("work_end_date", "")
    settings.setdefault("advance_amount", "")
    settings.setdefault("final_payment_amount", "")
    settings.setdefault("payments", [])
    settings.setdefault("working_group_text", "")
    settings.setdefault("intro_override", "")
    settings.setdefault("payments_override", "")
    settings.setdefault("communications_override", "")

    if not settings.get("customer_name") and counterparty:
        settings["customer_name"] = _counterparty_display_name(counterparty)
    if not settings.get("object_address"):
        settings["object_address"] = (counterparty.get("work_address") if counterparty else "") or project.get("address") or project.get("project_name") or ""
    if not settings.get("customer_email") and counterparty:
        settings["customer_email"] = counterparty.get("email") or ""
    if not settings.get("customer_phone") and counterparty:
        settings["customer_phone"] = counterparty.get("phone") or ""
    if not settings.get("passport_series_number") and counterparty:
        settings["passport_series_number"] = counterparty.get("passport_series_number") or ""
    if not settings.get("passport_issued_by") and counterparty:
        settings["passport_issued_by"] = counterparty.get("passport_issued_by") or ""
    if not settings.get("passport_department_code") and counterparty:
        settings["passport_department_code"] = counterparty.get("passport_department_code") or ""
    if not settings.get("registration_address") and counterparty:
        settings["registration_address"] = counterparty.get("registration_address") or counterparty.get("legal_address") or ""
    if not settings.get("price_total"):
        settings["price_total"] = estimate_total

    template_path = _get_template_path()
    replacements = build_contract_replacements(project, counterparty, settings, estimate_total)

    doc = replace_placeholders_in_docx(template_path, replacements)

    contracts_dir = _get_contracts_dir()
    project_name = project.get("project_name") or project.get("address") or f"Проект_{project_id}"
    safe_name = sanitize_filename(project_name, f"Проект_{project_id}")
    contract_number = settings.get("contract_number") or project.get("contract") or f"Договор_{project_id}"
    safe_number = sanitize_filename(contract_number, f"Договор_{project_id}")
    project_dir = contracts_dir / f"{int(project_id):04d}_{safe_name}"
    project_dir.mkdir(parents=True, exist_ok=True)

    output_filename = f"{safe_number}.docx"
    output_path = project_dir / output_filename
    doc.save(str(output_path))

    document = create_or_update_contract_document(project_id)
    relative_path = storage_relative_path(output_path)
    update_document_file_path(document["id"], relative_path)

    return {
        "document_id": document["id"],
        "project_id": project_id,
        "file_path": relative_path,
        "output_filename": output_filename,
    }


def _find_soffice() -> str:
    for cmd in ("soffice", "libreoffice"):
        path = shutil.which(cmd)
        if path:
            return path
    raise RuntimeError(
        "LibreOffice не найден. Установите libreoffice-writer для конвертации DOCX → PDF."
    )


def _build_soffice_cmd(soffice_path: str, docx_path: str, outdir: str, profile_dir: str) -> list[str]:
    return [
        soffice_path,
        "-env:UserInstallation=" + profile_dir,
        "--headless",
        "--norestore",
        "--nodefault",
        "--nofirststartwizard",
        "--nolockcheck",
        "--convert-to",
        "pdf",
        "--outdir",
        outdir,
        docx_path,
    ]


_CONVERSION_TIMEOUT = 60


def generate_contract_pdf(project_id: int) -> dict:
    from webapp.db import (
        fetch_project,
        update_document_pdf_path,
    )

    result = generate_contract_docx(project_id)
    document_id = result["document_id"]

    from webapp.db import fetch_document
    document = fetch_document(document_id)
    if not document:
        raise ValueError(f"Document {document_id} not found")

    docx_relative = document.get("file_path") or ""
    if not docx_relative:
        raise FileNotFoundError("DOCX договора не найден. Сначала сформируйте DOCX.")

    docx_absolute = resolve_storage_path(docx_relative)
    if not docx_absolute or not docx_absolute.exists():
        raise FileNotFoundError(f"DOCX файл не найден на диске: {docx_relative}")

    pdf_path = docx_absolute.with_suffix(".pdf")
    outdir = str(docx_absolute.parent)

    soffice_path = _find_soffice()

    with tempfile.TemporaryDirectory(prefix="dekorcrm-lo-") as profile_dir:
        profile_uri = Path(profile_dir).as_uri()
        cmd = _build_soffice_cmd(soffice_path, str(docx_absolute), outdir, profile_uri)

        try:
            proc = subprocess.run(
                cmd,
                timeout=_CONVERSION_TIMEOUT,
                capture_output=True,
                text=True,
            )
        except subprocess.TimeoutExpired:
            raise RuntimeError(
                f"Конвертация DOCX → PDF превысила таймаут ({_CONVERSION_TIMEOUT} сек)."
            )

        if proc.returncode != 0:
            raise RuntimeError(
                f"LibreOffice вернул ошибку (код {proc.returncode}): "
                f"stdout={proc.stdout[:500]}, stderr={proc.stderr[:500]}"
            )

    if not pdf_path.exists():
        raise FileNotFoundError(
            f"PDF файл не создан после конвертации. Ожидался: {pdf_path}"
        )

    pdf_relative = storage_relative_path(pdf_path)
    update_document_pdf_path(document_id, pdf_relative)

    return {
        "document_id": document_id,
        "project_id": project_id,
        "pdf_path": pdf_relative,
        "pdf_filename": pdf_path.name,
    }